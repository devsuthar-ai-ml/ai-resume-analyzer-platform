import io
import os

import pytest
from docx import Document
from fastapi.testclient import TestClient

os.environ["DATABASE_URL"] = "sqlite:///./test.db"
os.environ["SECRET_KEY"] = "test-secret"

from app.core.database import Base, SessionLocal, engine, get_db
from app.main import app


@pytest.fixture
def db_session():
    Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


@pytest.fixture
def client(db_session):
    def override_get_db():
        try:
            yield db_session
        finally:
            pass

    app.dependency_overrides[get_db] = override_get_db
    with TestClient(app) as test_client:
        yield test_client
    app.dependency_overrides.clear()


@pytest.fixture
def auth_headers(client):
    payload = {
        "email": "alice@example.com",
        "full_name": "Alice Doe",
        "password": "StrongPassword123",
    }
    response = client.post("/api/auth/register", json=payload)
    token = response.json()["access_token"]
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture
def sample_docx_bytes() -> bytes:
    doc = Document()
    doc.add_heading("Resume", level=1)
    doc.add_paragraph("Python developer with 5 years experience in FastAPI, Docker, and PostgreSQL.")
    doc.add_paragraph("Education: Bachelor of Technology in Computer Science")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
